import json

import openpyxl
import pytest

from bruv_app import demo
from bruv_app.cli import keep_awake
from bruv_app.export import export_project
from bruv_app.pipeline import (
    Pipeline,
    create_project,
    extra_species,
    load_extra_species,
    update_extra_species,
)
from bruv_app.detect import ColorBlobDetector

from .conftest import demo_settings, expected_maxn, ground_truth_species, label_like_perfect_human


def test_pipeline_builds_one_card_per_fish_visit(project):
    fish = project.all_fish()
    assert project.frame_count() == 24
    # 10 scene fish, but returning same-species fish within 10 s share a card.
    assert 4 <= len(fish) < 10
    for f in fish:
        guesses = json.loads(f["guesses"])
        assert len(guesses) == 3
        assert f["status"] == "pending"  # nothing is pre-selected


def test_tiny_fish_get_cards_of_their_own(project, tmp_path, synthetic_videos):
    assert len(project.all_fish()) < len(project.tracklet_ids())  # normally, pieces are joined
    tiny = create_project(tmp_path / "tiny", synthetic_videos, None, demo_settings(tiny_fish_px=1e9))
    Pipeline(tiny, classifier=demo.DemoClassifier()).run()  # every fish counts as tiny here
    assert len(tiny.all_fish()) == len(tiny.tracklet_ids())


def test_cards_are_pure(project):
    """No card mixes species (the tracker's appearance gate stops identity swaps)."""
    for f in project.all_fish():
        species = {ground_truth_species(d) for d in project.fish_detections(f["id"])}
        assert len(species - {None}) == 1, (f["id"], species)


def test_top_guess_matches_truth_on_synthetic_data(project):
    for f in project.all_fish():
        truth = ground_truth_species(project.fish_detections(f["id"])[0])
        assert json.loads(f["guesses"])[0][0] == truth


def test_end_to_end_maxn_matches_ground_truth(project, tmp_path):
    label_like_perfect_human(project)
    out = export_project(project, tmp_path / "export.xlsx")
    wb = openpyxl.load_workbook(out)
    maxn = {r[0]: r[1] for r in wb["MaxN"].iter_rows(min_row=2, values_only=True)}
    assert maxn == expected_maxn() == {demo.RED: 2, demo.YELLOW: 3, demo.GREEN: 1}

    first = {r[0]: r[6] for r in wb["MaxN"].iter_rows(min_row=2, values_only=True)}
    assert first == {demo.RED: 1.0, demo.YELLOW: 3.0, demo.GREEN: 21.0}

    ws = wb[demo.SHEET_RECORD]
    col = {c.value: c.column for c in ws[1]}
    records = [
        {h: ws.cell(r, c).value for h, c in col.items()} for r in range(3, ws.max_row + 1)
    ]
    assert {r["species"] for r in records} == {demo.RED, demo.YELLOW, demo.GREEN}
    assert all(r["clip_in"] in ("GX010042", "GX020042") for r in records)
    assert all(r["length_ind"] is None for r in records)
    # The yellow school: one sighting row, counted 3 at its peak.
    assert max(r["count"] for r in records if r["species"] == demo.YELLOW) == 3

    details = list(wb["app_details"].iter_rows(min_row=2, values_only=True))
    assert len(details) == len(project.all_fish())
    assert all(row[10] is True for row in details)  # user agreed with top-1 everywhere


def test_detection_resumes_after_interruption(tmp_path, synthetic_videos):
    template = demo.make_mock_template(tmp_path / "t.xlsx")
    project = create_project(tmp_path / "p", synthetic_videos, template, demo_settings(batch_size=4))

    class Crashy(ColorBlobDetector):
        calls = 0

        def detect(self, images):
            Crashy.calls += 1
            if Crashy.calls == 3:
                raise KeyboardInterrupt
            return super().detect(images)

    with pytest.raises(KeyboardInterrupt):
        Pipeline(project, detector=Crashy(), classifier=demo.DemoClassifier()).run()
    assert project.frame_count() == 8
    Pipeline(project, detector=ColorBlobDetector(), classifier=demo.DemoClassifier()).run()
    times = [r["t"] for r in project.db.execute("SELECT t FROM frames ORDER BY seq")]
    assert len(times) == 24
    assert [round(t) for t in times] == list(range(24))


def test_rerun_keeps_reviews_and_refuses_to_wipe_them(project):
    first = project.all_fish()[0]["id"]
    project.label_fish(first, "species", species=demo.RED)
    Pipeline(project, classifier=demo.DemoClassifier()).run()  # all stages done: no-op
    assert project.fish(first)["species"] == demo.RED
    for stage in ("merge", "classify"):
        project.mark_stage(stage, False)
    with pytest.raises(RuntimeError):
        Pipeline(project, classifier=demo.DemoClassifier()).run()


def test_cards_built_without_the_species_model_are_rebuilt_on_the_next_run(project):
    """A run that fell back to colour histograms (BioCLIP 2 couldn't load) is redone with
    the model next time, but never once someone has started reviewing."""
    from bruv_app.classify import HistogramClassifier

    Pipeline(project, classifier=HistogramClassifier()).run(force=True)
    assert project.get_meta("classifier") == "hsv-histogram"
    settings = project.load_settings()
    settings.classifier = "bioclip"
    project.save_settings(settings)
    Pipeline(project, classifier=demo.DemoClassifier()).run()
    assert project.get_meta("classifier") == "demo-colour"
    assert all(json.loads(f["guesses"]) for f in project.all_fish())

    Pipeline(project, classifier=HistogramClassifier()).run(force=True)
    project.label_fish(project.all_fish()[0]["id"], "species", species=demo.RED)
    Pipeline(project, classifier=demo.DemoClassifier()).run()
    assert project.get_meta("classifier") == "hsv-histogram"  # reviews are never wiped


def test_extra_names_are_guessed_picked_and_exported_as_other(tmp_path, synthetic_videos, dependent_template):
    extra = ["batfish", "triggerfish", "bluespot mullet"]
    project = create_project(tmp_path / "p", synthetic_videos, dependent_template, demo_settings(), extra)
    assert extra_species(project) == ["batfish", "triggerfish"]  # a template name isn't an extra
    Pipeline(project, classifier=demo.DemoClassifier()).run()
    assert project.get_meta("guess_labels")[-2:] == ["batfish", "triggerfish"]

    a, b, c = project.all_fish()[:3]
    project.label_fish(a["id"], "species", species="batfish")
    project.label_fish(b["id"], "species", species="triggerfish", comment="near the bait")
    project.label_fish(c["id"], "species", species="bluespot mullet")
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "out.xlsx"))
    ws = wb[demo.SHEET_RECORD]
    col = {cell.value: cell.column for cell in ws[1]}
    rows = [{h: ws.cell(r, c).value for h, c in col.items()} for r in range(2, ws.max_row + 1)]
    by_comment = {r["comment"]: r for r in rows}
    # Extra names stay inside the template's dropdown: its catch-all, with the name in the comment.
    assert (by_comment["batfish"]["species"], by_comment["batfish"]["type"]) == ("other (indicate in comment)", "interest")
    assert by_comment["triggerfish; near the bait"]["species"] == "other (indicate in comment)"
    assert any(r["species"] == "bluespot mullet" and r["type"] == "mullet" for r in rows)
    # MaxN counts each name on its own instead of lumping them under "other".
    maxn = {r[0] for r in wb["MaxN"].iter_rows(min_row=2, values_only=True)}
    assert {"batfish", "triggerfish", "bluespot mullet"} <= maxn


def test_names_dropped_from_the_lists_still_export_as_other(tmp_path, synthetic_videos, dependent_template):
    """A card recorded with a name later renamed in the extra list stays valid in Excel."""
    project = create_project(tmp_path / "p", synthetic_videos, dependent_template, demo_settings(), ["batfish"])
    Pipeline(project, classifier=demo.DemoClassifier()).run()
    first = project.all_fish()[0]["id"]
    project.label_fish(first, "species", species="batfish")
    update_extra_species(project, ["batfish / spadefish"])
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "out.xlsx"))
    ws = wb[demo.SHEET_RECORD]
    col = {cell.value: cell.column for cell in ws[1]}
    assert (ws.cell(2, col["species"]).value, ws.cell(2, col["comment"]).value) == ("other (indicate in comment)", "batfish")


def test_export_has_the_corrections_log(project, tmp_path):
    f = next(f for f in project.all_fish() if len(project.fish_tracklets(f["id"])) > 1)
    project.unmerge_fish(f["id"])
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "out.xlsx"))
    rows = list(wb["corrections"].iter_rows(values_only=True))
    assert rows[0] == ("at", "action", "card", "details")
    assert rows[1][1:3] == ("multiple_fish", f["id"])


def test_extra_species_file_next_to_the_template_wins(tmp_path):
    default = load_extra_species(None)
    assert len(default) == 33 and {"triggerfish", "fusilier", "moorish idol", "batfish / spadefish"} <= set(default)
    assert not {"salmon", "cod", "batfish"} & set(default)
    (tmp_path / "extra_species.txt").write_text("# my region\nkelpfish\n\nkelpfish\nsea perch\n")
    assert load_extra_species(tmp_path / "template.xlsx") == ["kelpfish", "sea perch"]


def test_embedding_resumes_after_a_stop(tmp_path, synthetic_videos, monkeypatch):
    monkeypatch.setattr(Pipeline, "EMBED_CHUNK", 8)

    class Counting(demo.DemoClassifier):
        def __init__(self, stop_after=None):
            self.calls, self.crops, self.stop_after = 0, 0, stop_after

        def embed(self, crops):
            self.calls += 1
            if self.stop_after is not None and self.calls > self.stop_after:
                raise KeyboardInterrupt
            self.crops += len(crops)
            return super().embed(crops)

    template = demo.make_mock_template(tmp_path / "t.xlsx")
    project = create_project(tmp_path / "p", synthetic_videos, template, demo_settings())
    first = Counting(stop_after=2)
    with pytest.raises(KeyboardInterrupt):
        Pipeline(project, classifier=first).run()
    assert first.crops == 16  # two chunks done and saved before the stop
    second = Counting()
    Pipeline(project, classifier=second).run()
    embedded = project.db.execute("SELECT COUNT(*) FROM detections WHERE embedding IS NOT NULL").fetchone()[0]
    assert second.crops == embedded - 16  # nothing embedded twice

    fresh = create_project(tmp_path / "q", synthetic_videos, template, demo_settings())
    Pipeline(fresh, classifier=demo.DemoClassifier()).run()
    assert [json.loads(f["guesses"]) for f in project.all_fish()] == [json.loads(f["guesses"]) for f in fresh.all_fish()]


def test_new_species_names_resume_and_keep_reviews(project, monkeypatch):
    """New extra names only recompute the guesses: cards and reviews stay. The names are
    saved as they're encoded, so a stopped run carries on with the rest."""
    monkeypatch.setattr(Pipeline, "LABEL_CHUNK", 2)

    class CountingText(demo.DemoClassifier):
        def __init__(self, stop_after=None):
            self.encoded, self.calls, self.stop_after = [], 0, stop_after

        def text_embeddings(self, labels):
            self.calls += 1
            if self.stop_after is not None and self.calls > self.stop_after:
                raise KeyboardInterrupt
            self.encoded += labels
            return super().text_embeddings(labels)

    reviewed = project.all_fish()[0]["id"]
    project.label_fish(reviewed, "species", species=demo.RED)
    cards = [f["id"] for f in project.all_fish()]
    update_extra_species(project, ["batfish", "goby", "wrasse"])
    first = CountingText(stop_after=1)
    with pytest.raises(KeyboardInterrupt):
        Pipeline(project, classifier=first).run()
    assert first.encoded == ["batfish", "goby"]  # the template's names came from the saved ones
    second = CountingText()
    Pipeline(project, classifier=second).run()
    assert second.encoded == ["wrasse"]
    assert project.get_meta("guess_labels")[-3:] == ["batfish", "goby", "wrasse"]
    assert [f["id"] for f in project.all_fish()] == cards
    assert project.fish(reviewed)["species"] == demo.RED


def test_keep_awake_is_safe():
    with keep_awake():
        pass


def test_clip_names_follow_the_template_clip_list(tmp_path, synthetic_videos, dependent_template):
    project = create_project(tmp_path / "p", synthetic_videos, dependent_template, demo_settings())
    assert project.get_meta("dependent") == {"species": "type"}
    Pipeline(project, classifier=demo.DemoClassifier()).ingest()
    # The template's clip_in dropdown lists clip_1, clip_2, ... so use those, in chapter order.
    assert [c["name"] for c in project.chapters()] == ["clip_1", "clip_2"]


def test_analysis_window_limits_frames(tmp_path, synthetic_videos):
    template = demo.make_mock_template(tmp_path / "t.xlsx")
    settings = demo_settings(start_s=5.0, duration_s=10.0, sample_fps=2.0)
    project = create_project(tmp_path / "p", synthetic_videos, template, settings)
    Pipeline(project, classifier=demo.DemoClassifier()).run()
    times = [r["t"] for r in project.db.execute("SELECT t FROM frames ORDER BY seq")]
    assert len(times) == 20 and times[0] == pytest.approx(5.0) and times[-1] < 15.0
