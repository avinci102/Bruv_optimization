import hashlib
from datetime import time

import openpyxl
import pytest

from bruv_app import demo
from bruv_app.template import export_workbook, read_raw_validations, read_template, seconds_to_time


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_reads_dropdowns_from_all_validation_kinds(mock_template):
    info = read_template(mock_template)
    assert info.record_sheet == demo.SHEET_RECORD
    assert info.header_row == 1
    assert info.next_row == 3  # the example row 2 is left alone
    assert info.options("species")[:3] == [demo.RED, demo.YELLOW, demo.GREEN]  # x14 cross-sheet
    assert info.options("behavior") == demo.BEHAVIORS  # x14 cross-sheet
    assert info.options("type") == demo.TYPES  # named range
    assert info.options("sex") == demo.SEXES  # literal list
    assert sum(v.from_extension for v in info.validations) == 2


def test_reads_a_dependent_species_list(dependent_template):
    info = read_template(dependent_template)
    assert info.options("type") == ["mullet", "goatfish"]
    # species =INDIRECT(E2): every type's named range, in order.
    assert info.options("species") == ["mullet family", "bluespot mullet", "goatfish family", "yellowfin goatfish"]
    parent, groups = info.dependent["species"]
    assert parent == "type" and groups["yellowfin goatfish"] == "goatfish"
    # One x14 dropdown covers both clip columns.
    assert info.options("clip_in") == info.options("clip_out") == ["clip_1", "clip_2", "clip_3"]


def test_export_fills_type_from_species_and_keeps_two_range_dropdown(dependent_template, tmp_path):
    info = read_template(dependent_template)
    out = tmp_path / "out.xlsx"
    rows = [
        {"clip_in": "clip_1", "clip_out": "clip_1", "type": None, "species": "bluespot mullet", "count": 1},
        {"clip_in": "clip_2", "clip_out": "clip_2", "type": None, "species": None, "count": 1},
    ]
    export_workbook(info, out, rows, ["species"], [], ["fish_id"], [])
    ws = openpyxl.load_workbook(out)[demo.SHEET_RECORD]
    assert (ws["E2"].value, ws["F2"].value) == ("mullet", "bluespot mullet")
    assert ws["E3"].value is None  # unknown fish: no species, so no type
    again = read_template(out)
    assert again.options("clip_in") == again.options("clip_out") == ["clip_1", "clip_2", "clip_3"]
    assert again.options("species") == info.options("species")


def test_falls_back_to_dropdown_sheet_column_when_no_validation(tmp_path):
    wb = openpyxl.Workbook()
    dd = wb.active
    dd.title = demo.SHEET_DROPDOWN
    dd.append(["species", "sex"])
    dd.append(["Fish one", "M"])
    dd.append(["Fish two", "F"])
    rec = wb.create_sheet(demo.SHEET_RECORD)
    rec.append(["clip_in", "time_in", "species", "sex"])
    path = tmp_path / "plain.xlsx"
    wb.save(path)
    info = read_template(path)
    assert info.options("species") == ["Fish one", "Fish two"]
    assert info.options("sex") == ["M", "F"]


def test_finds_header_row_below_a_title(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append(["My BRUV sheet"])
    ws.append(["clip_in", "time_in", "species"])
    path = tmp_path / "titled.xlsx"
    wb.save(path)
    info = read_template(path)
    assert info.header_row == 2 and info.headers["species"] == 3


def test_export_writes_by_header_name_into_a_copy(mock_template, tmp_path):
    before = sha(mock_template)
    info = read_template(mock_template)
    out = tmp_path / "out.xlsx"
    rows = [
        {"clip_in": "GX010042", "time_in": seconds_to_time(75), "clip_out": "GX010042",
         "time_out": seconds_to_time(80), "species": demo.RED, "count": 2, "behavior": "feeding",
         "not_a_header": "x"},
    ]
    result = export_workbook(info, out, rows, ["species", "maxn"], [[demo.RED, 2]], ["fish_id"], [[1]])
    assert sha(mock_template) == before, "the original template must never change"
    assert result["first_row"] == 3 and result["skipped_headers"] == ["not_a_header"]

    wb = openpyxl.load_workbook(out)
    ws = wb[demo.SHEET_RECORD]
    col = {c.value: c.column for c in ws[1]}
    assert ws.cell(3, col["species"]).value == demo.RED
    assert ws.cell(3, col["time_in"]).value == time(0, 1, 15)
    assert ws.cell(3, col["time_in"]).number_format == "hh:mm:ss"
    assert ws.cell(3, col["count"]).value == 2
    assert ws.cell(3, col["length_ind"]).value is None
    assert ws.cell(2, col["clip_in"]).value == "clip_1"  # example row untouched
    for letter in demo.HIDDEN:
        assert ws.column_dimensions[letter].hidden
        assert ws[f"{letter}3"].value is None
    assert wb["MaxN"]["B2"].value == 2
    assert wb["app_details"]["A2"].value == 1


def test_export_keeps_every_dropdown(mock_template, tmp_path):
    info = read_template(mock_template)
    out = tmp_path / "out.xlsx"
    export_workbook(info, out, [], ["species"], [], ["fish_id"], [])
    formulas = {v.formula1 for v in read_raw_validations(out, demo.SHEET_RECORD)}
    assert formulas == {v.formula1 for v in info.validations}
    # The exported copy reads back with the same lists.
    assert read_template(out).dropdowns == info.dropdowns


def test_refuses_to_overwrite_the_template(mock_template):
    info = read_template(mock_template)
    with pytest.raises(ValueError):
        export_workbook(info, mock_template, [], [], [], [], [])


def test_seconds_to_time():
    assert seconds_to_time(3725.4) == time(1, 2, 5)
    assert seconds_to_time(-3) == time(0, 0, 0)


def test_reattached_dropdowns_keep_their_error_settings(mock_template, tmp_path):
    info = read_template(mock_template)
    out = tmp_path / "out.xlsx"
    export_workbook(info, out, [], ["species"], [], ["fish_id"], [])
    species = next(v for v in read_raw_validations(out, demo.SHEET_RECORD) if v.sqref == "N2:N1000")
    # Typed values outside the list must still be rejected in Excel.
    assert species.options.get("showErrorMessage") is True
    assert species.options.get("showInputMessage") is True


def test_export_reports_template_features_it_cannot_keep(mock_template, tmp_path, caplog, monkeypatch):
    import logging
    import warnings

    info = read_template(mock_template)
    real = openpyxl.load_workbook

    def noisy(*args, **kwargs):
        warnings.warn("Conditional Formatting extension is not supported and will be removed")
        return real(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", noisy)
    with caplog.at_level(logging.WARNING):
        export_workbook(info, tmp_path / "out.xlsx", [], ["species"], [], ["fish_id"], [])
    assert any("Conditional Formatting" in r.message for r in caplog.records)
    assert not any("Data Validation" in r.message for r in caplog.records)


def test_dropdowns_on_other_sheets_survive_export(tmp_path):
    template = demo.make_mock_template(tmp_path / "base.xlsx")
    # Give the info sheet a cross-sheet (x14) dropdown too.
    path = demo.inject_x14_validations(
        template.read_bytes(),
        tmp_path / "with_info_dropdown.xlsx",
        {demo.SHEET_INFO: [(f"'{demo.SHEET_DROPDOWN}'!$C$2:$C$4", "B3")]},
    )
    assert any(v.from_extension for v in read_raw_validations(path, demo.SHEET_INFO))
    out = tmp_path / "out.xlsx"
    export_workbook(read_template(path), out, [], ["species"], [], ["fish_id"], [])
    kept = read_raw_validations(out, demo.SHEET_INFO)
    assert [(v.formula1, v.sqref) for v in kept] == [(f"'{demo.SHEET_DROPDOWN}'!$C$2:$C$4", "B3")]
    assert len(read_raw_validations(out, demo.SHEET_RECORD)) == 4


def test_external_link_with_alternate_paths_stays_valid(mock_template, tmp_path):
    """A link saved with two paths, pointing at the second: openpyxl keeps the first, and
    Excel refuses a file whose link points at a relationship that isn't there."""
    import re
    import zipfile

    path_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath"
    parts = {
        "xl/externalLinks/externalLink1.xml": (
            '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<externalBook r:id="rId1"><sheetNames><sheetName val="info"/></sheetNames></externalBook>'
            "</externalLink>"
        ),
        "xl/externalLinks/_rels/externalLink1.xml.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId2" Type="{path_type}" Target="../old/template.xlsx" TargetMode="External"/>'
            f'<Relationship Id="rId1" Type="{path_type}" Target="/Users/x/old/template.xlsx" TargetMode="External"/>'
            "</Relationships>"
        ),
    }
    linked = tmp_path / "linked.xlsx"
    with zipfile.ZipFile(mock_template) as src, zipfile.ZipFile(linked, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename).decode()
            if item.filename == "xl/workbook.xml":
                data = data.replace("</sheets>", '</sheets><externalReferences><externalReference r:id="rIdLink"/></externalReferences>')
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace("</Relationships>", (
                    '<Relationship Id="rIdLink" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/externalLink" Target="/xl/externalLinks/externalLink1.xml"/></Relationships>'))
            elif item.filename == "[Content_Types].xml":
                data = data.replace("</Types>", (
                    '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/'
                    'vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>'))
            dst.writestr(item, data)
        for name, xml in parts.items():
            dst.writestr(name, xml)

    out = tmp_path / "out.xlsx"
    export_workbook(read_template(linked), out, [], ["species"], [], ["fish_id"], [])
    with zipfile.ZipFile(out) as zf:
        book_id = re.search(r'externalBook[^>]*r:id="([^"]+)"', zf.read("xl/externalLinks/externalLink1.xml").decode()).group(1)
        rels = zf.read("xl/externalLinks/_rels/externalLink1.xml.rels").decode()
    assert f'Id="{book_id}"' in rels
    assert 'Target="/Users/x/old/template.xlsx"' in rels  # the path the link pointed at


@pytest.mark.filterwarnings("ignore:Data Validation extension")  # the setup's own re-save
def test_export_keeps_rich_text_and_gives_new_cells_their_column_format(tmp_path):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment

    path = demo.make_mock_template(tmp_path / "plain.xlsx")
    wb = openpyxl.load_workbook(path)
    # The italic run holding only a space must keep it (and keep the file valid for Excel).
    wb[demo.SHEET_INFO]["C1"] = CellRichText(
        "Read ", TextBlock(InlineFont(b=True), "this"), TextBlock(InlineFont(i=True), " "), "first"
    )
    count = next(c.column_letter for c in wb[demo.SHEET_RECORD][1] if c.value == "count")
    wb[demo.SHEET_RECORD].column_dimensions[count].alignment = Alignment(horizontal="center")
    styled = tmp_path / "styled.xlsx"
    wb.save(styled)

    out = tmp_path / "out.xlsx"
    export_workbook(read_template(styled), out, [{"species": demo.RED, "count": 3}], ["species"], [], ["fish_id"], [])
    wb = openpyxl.load_workbook(out, rich_text=True)
    text = wb[demo.SHEET_INFO]["C1"].value
    assert isinstance(text, CellRichText) and any(getattr(part, "font", None) and part.font.b for part in text)
    assert str(text) == "Read this first"
    new = wb[demo.SHEET_RECORD][f"{count}3"]
    assert new.value == 3 and new.alignment.horizontal == "center"


def test_export_warns_about_shapes_it_cannot_keep(mock_template, tmp_path, caplog):
    import logging
    import zipfile

    path = tmp_path / "with_shape.xlsx"
    with zipfile.ZipFile(mock_template) as src, zipfile.ZipFile(path, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/drawings/drawing1.xml", "<xdr:wsDr><xdr:twoCellAnchor><xdr:sp><xdr:txBody/></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>")
    with caplog.at_level(logging.WARNING):
        export_workbook(read_template(path), tmp_path / "out.xlsx", [], ["species"], [], ["fish_id"], [])
    assert any("shapes or text boxes" in r.message for r in caplog.records)
