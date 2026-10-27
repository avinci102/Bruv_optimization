"""Wave-1 rules: one record row per card, MaxN evidence frames, tracker options and trials."""

from pathlib import Path

import numpy as np
import openpyxl
import pytest

from bruv_app import demo
from bruv_app.cli import main
from bruv_app.detect import Detection
from bruv_app.export import export_project
from bruv_app.follow import mark, start
from bruv_app.metrics import FishLabel, card_sightings
from bruv_app.track import track_all
from bruv_app.trials import BASELINE, apply_changes, cautions, compare, open_trial, table, trial_names, try_rules

from .conftest import label_like_perfect_human


def box(cx, cy, score=0.9, w=0.1, h=0.05):
    return Detection(cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2, score, None)


def test_each_answered_card_is_a_row_of_one():
    fish = [FishLabel(1, "red", 0.0, 4.0), FishLabel(2, "red", 5.0, 9.0), FishLabel(3, "blue", 1.0, 2.0)]
    rows = card_sightings(fish)
    assert [(s.species, s.fish_ids, s.count) for s in rows] == [("red", [1], 1), ("blue", [3], 1), ("red", [2], 1)]


def _records(path):
    ws = openpyxl.load_workbook(path)[demo.SHEET_RECORD]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=3, values_only=True) if any(r)]  # row 2: an example


def test_the_export_gives_each_card_a_row_and_shows_the_maxn_frame(project, tmp_path):
    label_like_perfect_human(project)
    answered = [f for f in project.all_fish() if f["status"] in ("species", "unknown")]
    out = export_project(project, tmp_path / "cards.xlsx")
    rows = _records(out)
    assert len(rows) == len(answered) and {r["count"] for r in rows} == {1}
    wb = openpyxl.load_workbook(out)
    status = {r[0]: r[1] for r in wb["review_status"].iter_rows(values_only=True)}
    assert status["record rows"].startswith("one per answered card")
    maxn = list(wb["MaxN"].iter_rows(values_only=True))
    assert maxn[0][-1] == "evidence_frame"
    for row in maxn[1:]:
        assert row[-1] and Path(row[-1]).exists() and Path(row[-1]).parent.name == "cards_maxn"
    assert all(r[-1] for r in list(wb["MaxN_by_family"].iter_rows(values_only=True))[1:])

    # The rows from before: same species within sighting_gap_s share one.
    settings = project.load_settings()
    settings.record_rows = "grouped"
    project.save_settings(settings)
    grouped = _records(export_project(project, tmp_path / "grouped.xlsx"))
    assert len(grouped) < len(rows) and max(r["count"] for r in grouped) > 1


def test_faint_boxes_carry_a_track_but_never_start_one():
    # The fish fades below the threshold in frame 2; a faint box elsewhere is just noise.
    frames = [[(0, box(0.20, 0.5))], [(1, box(0.25, 0.5))], [], [(3, box(0.35, 0.5))]]
    faint = {2: [(2, box(0.30, 0.5, 0.2)), (9, box(0.8, 0.2, 0.2))]}
    from bruv_app.track import Tracker  # noqa: PLC0415

    for use in (False, True):
        tracker = Tracker(max_missed=0)
        for i, dets in enumerate(frames):
            tracker.update(i, dets, faint.get(i) if use else None)
        tracks = tracker.finish()
        keys = sorted([k for _, k in t.members] for t in tracks)
        assert keys == ([[0, 1, 2, 3]] if use else [[0, 1], [3]])
    # A faint box never takes a strong box's place on a track.
    tracker = Tracker()
    tracker.update(0, [(0, box(0.2, 0.5))])
    tracker.update(1, [(1, box(0.25, 0.5))], [(2, box(0.24, 0.5, 0.2))])
    assert [k for _, k in tracker.finish()[0].members] == [0, 1]


def test_smoothing_keeps_some_of_the_earlier_course():
    frames = [[(0, box(0.10, 0.5))], [(1, box(0.20, 0.5))], [(2, box(0.24, 0.5))]]
    plain = track_all(frames)[0].velocity
    smooth = track_all(frames, velocity_smoothing=0.5)[0].velocity
    assert plain[0] == pytest.approx(0.04) and smooth[0] == pytest.approx(0.07)


def test_settings_are_read_as_their_types():
    s = apply_changes(demo_settings_like(), {"track_faint": "true", "track_velocity_smoothing": "0.5", "tiny_fish_px": "30"})
    assert s.track_faint is True and s.track_velocity_smoothing == 0.5 and s.tiny_fish_px == 30.0
    with pytest.raises(ValueError, match="No setting"):
        apply_changes(demo_settings_like(), {"nonsense": "1"})
    with pytest.raises(ValueError):
        apply_changes(demo_settings_like(), {"track_faint": "maybe"})


def test_a_trial_says_when_faint_boxes_were_never_kept(project):
    threshold = project.load_settings().det_threshold
    faint = project.db.execute("SELECT COUNT(*) FROM detections WHERE score < ?", (threshold,)).fetchone()[0]
    assert bool(cautions(project, {"track_faint": "true"})) == (faint == 0)
    project.db.execute("DELETE FROM detections WHERE score < ? AND tracklet_id IS NULL", (threshold,))
    assert "processed before" in cautions(project, {"track_faint": "true"})[0]
    assert cautions(project, {"track_velocity_smoothing": "0.5"}) == []


def demo_settings_like():
    from bruv_app.config import Settings  # noqa: PLC0415

    return Settings()


def test_a_trial_rebuilds_a_copy_and_is_scored_against_the_project(project, capsys):
    label_like_perfect_human(project)
    card = next(f["id"] for f in project.all_fish() if len(project.fish_detections(f["id"])) >= 3)
    follow_id = start(project, card)
    for d in project.fish_detections(card):
        mark(project, follow_id, d["frame_id"], "box", det_id=d["id"])
    cards, settings_before = len(project.all_fish()), project.settings_path.read_text()

    main(["try-rules", "--project", str(project.root), "--name", "faint", "--set", "track_faint=true",
          "track_velocity_smoothing=0.5"])
    out = capsys.readouterr().out
    assert "baseline" in out and "faint" in out
    assert trial_names(project) == [BASELINE, "faint"]
    # The project itself is untouched; the trial has its own settings.
    assert len(project.all_fish()) == cards and project.settings_path.read_text() == settings_before
    trial = open_trial(project, "faint")
    assert trial.load_settings().track_faint is True and trial.settings_path.name == "faint.settings.json"
    assert not (project.root / "snapshots").exists() or not any("rebuilt" in p.name for p in (project.root / "snapshots").iterdir())

    scores = compare(project)
    assert [s.name for s in scores] == [BASELINE, "faint"]
    base = scores[0]
    assert base.changes == {} and base.checked > 0 and base.followed == 1 and base.follow_found == base.follow_present
    lines = table(scores)
    assert lines[0].startswith("trial") and any("track_faint" in line for line in lines)
    main(["compare-rules", "--project", str(project.root)])
    assert "followed boxed" in capsys.readouterr().out
    with pytest.raises(ValueError, match="No trial"):
        open_trial(project, "nothing")
    with pytest.raises(ValueError, match="name"):
        try_rules(project, "bad name!", {})
