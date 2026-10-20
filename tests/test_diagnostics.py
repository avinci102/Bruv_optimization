"""Continuity diagnostics: they record evidence and warn the reviewer, never change cards."""

import json

import numpy as np
import openpyxl
import pytest
from fastapi.testclient import TestClient

from bruv_app import demo
from bruv_app.continuity import Edge, TrackletSummary, merge_tracklets
from bruv_app.detect import Detection
from bruv_app.diagnostics import card_flags, pieces
from bruv_app.export import export_project
from bruv_app.server import create_app
from bruv_app.track import Tracker, missed_frames, track_all

from .conftest import label_like_perfect_human

RED, BLUE = np.array([1.0, 0.0]), np.array([0.0, 1.0])


def _fish(x, y=0.4, app=RED):
    return Detection(x, y, x + 0.1, y + 0.1, 0.9, app)


def test_tracker_records_each_link_without_changing_it():
    lone = track_all([[(1, _fish(0.1))], [(2, _fish(0.12))], [(3, _fish(0.14))]])
    assert len(lone) == 1 and set(lone[0].links) == {2, 3}  # every box after the first
    cost, margin, history = lone[0].links[3]
    assert cost >= 0 and margin is None and history == pytest.approx(1.0)  # no other option

    # Two fish close together: each link had another option, so it has a margin.
    pair = track_all([
        [(1, _fish(0.1)), (2, _fish(0.25, app=BLUE))],
        [(3, _fish(0.12)), (4, _fish(0.27, app=BLUE))],
    ], min_appearance=0.0)
    links = {k: v for t in pair for k, v in t.links.items()}
    assert set(links) == {3, 4} and all(m is not None and m > 0 for _, m, _ in links.values())
    assert [t.members for t in pair] == [[(0, 1), (1, 3)], [(0, 2), (1, 4)]]


@pytest.mark.parametrize(
    ("fps", "seconds", "frames", "bridged_s"),
    [(4.0, 1.0, 4, 1.25), (3.0, 1.0, 3, 4 / 3), (1.0, 1.0, 1, 2.0), (2.5, 1.0, 2, 1.2), (4.0, 0.0, 1, 0.5), (4.0, 0.5, 2, 0.75)],
)
def test_missed_time_boundaries(fps, seconds, frames, bridged_s):
    """The tracker's allowance for a fish going unseen, as built (review V2 asked for these
    boundaries to be pinned down before anyone changes them)."""
    m = missed_frames(seconds, fps)
    assert m == frames and (m + 1) / fps == pytest.approx(bridged_s)
    for unseen, same_track in ((m - 1, True), (m, True), (m + 1, False)):
        frames_seen = [[(1, _fish(0.1))]] + [[] for _ in range(unseen)] + [[(2, _fish(0.1))]]
        tracks = track_all(frames_seen, max_missed=m)
        assert (len(tracks) == 1) == same_track, (fps, seconds, unseen)


def _summary(i, start, end, emb, colour=None, size=0.1):
    emb = np.asarray(emb, float) / np.linalg.norm(emb)
    return TrackletSummary(i, start, end, emb, size, size, 100.0, None if colour is None else np.asarray(colour, float))


def test_merge_keeps_every_join_it_weighed():
    pieces_ = [
        _summary(1, 0, 1, [1, 0, 0], colour=[1, 0]),
        _summary(2, 2, 3, [1, 0.05, 0], colour=[1, 0]),  # most like 1: joined
        _summary(3, 2.5, 3.5, [1, 0.1, 0]),  # also like 1, but 1 already has its successor: lost
        _summary(4, 4, 5, [1, 0, 0], colour=[0, 1]),  # colours unlike 1's and 2's
        _summary(5, 4, 5, [1, 0, 0], size=0.5),  # five times the size
        _summary(6, 4.5, 5, [0, 1, 0]),  # looks nothing alike
    ]
    evidence: list[Edge] = []
    chains = merge_tracklets(pieces_, 10, 0.85, 3.0, 0, 0.7, evidence)
    assert chains == merge_tracklets(pieces_, 10, 0.85, 3.0, 0, 0.7)  # the record changes nothing
    decided = {(e.a, e.b): e.decision for e in evidence}
    assert decided[(1, 2)] == "joined" and decided[(1, 3)] == "lost"
    assert decided[(2, 4)] == "colour" and decided[(2, 5)] == "size" and decided[(2, 6)] == "appearance"
    edge = next(e for e in evidence if (e.a, e.b) == (1, 2))
    assert edge.gap == 1 and edge.similarity == pytest.approx(float(np.dot(pieces_[0].embedding, pieces_[1].embedding)))


def test_pipeline_stores_link_and_merge_evidence(project):
    rows = project.db.execute("SELECT tracklet_id, link_cost, history_sim FROM detections ORDER BY id").fetchall()
    first = {}
    for r in rows:
        first.setdefault(r["tracklet_id"], r)
    continuations = [r for r in rows if first[r["tracklet_id"]] is not r]
    assert continuations and all(r["link_cost"] is not None for r in continuations)
    assert all(first[t]["link_cost"] is None for t in first)  # a track's first box has no link
    edges = project.merge_edges()
    joined = {(e["a"], e["b"]) for e in edges if e["decision"] == "joined"}
    chained = {(a, b) for f in project.all_fish() for a, b in zip(project.fish_tracklets(f["id"]), project.fish_tracklets(f["id"])[1:])}
    assert joined == chained  # every join on a card is in the evidence, and nothing else was joined
    info = json.loads(project.runs("merge")[-1]["info"])
    assert info["decisions"]["joined"] == len(joined) and info["joins_weighed"] == len(edges)


def _bad_merge(project):
    """Fold a yellow card into a red one that's on screen at the same time."""
    label_like_perfect_human(project)
    cards = [f for f in project.all_fish() if f["status"] == "species"]
    red = next(f for f in cards if f["species"] == demo.RED and f["end_t"] > 17)
    yellow = next(f for f in cards if f["species"] == demo.YELLOW and f["start_t"] < 17 < f["end_t"])
    project.merge_fish(red["id"], yellow["id"])
    return red["id"]


def test_flags_on_a_card_that_mixes_fish(project):
    good = next(f for f in project.all_fish() if len(project.fish_tracklets(f["id"])) > 1)
    assert card_flags(project, good["id"], 0.85) == []  # pieces of one fish: no warnings
    card = _bad_merge(project)
    parts = pieces(project, card)
    assert len(parts) >= 2 and all(p.guesses for p in parts)
    kinds = [f.kind for f in card_flags(project, card, 0.85)]
    assert kinds == ["pieces look different", "pieces disagree on species", "two boxes at once"]


def test_a_flagged_card_needs_confirming(project, tmp_path):
    card = _bad_merge(project)
    client = TestClient(create_app(project.root))
    detail = client.get(f"/api/fish/{card}").json()
    assert len(detail["flags"]) == 3 and len(detail["pieces"]) >= 2
    assert all(p["guesses"] and p["best_det"] for p in detail["pieces"])
    answer = {"status": "species", "species": demo.RED}
    refused = client.post(f"/api/fish/{card}/label", json=answer)
    assert refused.status_code == 409 and "Check this card first" in refused.json()["detail"]
    assert project.fish(card)["status"] == "pending"
    assert client.post(f"/api/fish/{card}/label", json={**answer, "confirm": True}).status_code == 200
    review = json.loads(project.last_review(card)["details"])
    assert review == {"flags": ["pieces look different", "pieces disagree on species", "two boxes at once"], "confirmed": True}
    # Skipping a flagged card needs no confirming.
    assert client.post(f"/api/fish/{card}/label", json={"status": "skipped"}).status_code == 200

    client.post(f"/api/fish/{card}/label", json={**answer, "confirm": True})
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "out.xlsx"))
    details = {r[0]: r for r in wb["app_details"].iter_rows(min_row=2, values_only=True)}
    assert details[card][-1] == "pieces look different; pieces disagree on species; two boxes at once (confirmed one fish)"
    status = dict(list(wb["review_status"].iter_rows(values_only=True))[1:])
    assert status["cards with warnings (may hold more than one fish)"].startswith("1: 1 answered after confirming")


def test_card_detail_lists_the_trackers_closest_call(project):
    client = TestClient(create_app(project.root))
    card = next(f for f in project.all_fish() if len(project.fish_detections(f["id"])) > 2)
    detail = client.get(f"/api/fish/{card['id']}").json()
    assert detail["flags"] == []
    for call in detail["closest_calls"]:
        assert {"det_id", "frame_id", "clip", "local_t", "box", "margin"} <= set(call)


def test_tracker_without_colour_or_direction_still_records(tmp_path):
    tracker = Tracker(direction=False)
    tracker.update(0, [(1, _fish(0.1, app=None))])
    tracker.update(1, [(2, _fish(0.11, app=None))])
    (track,) = tracker.finish()
    assert track.links[2][2] is None  # no colours, no history likeness
