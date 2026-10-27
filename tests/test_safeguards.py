"""The review's mandatory scientific safeguards (review V2, section F)."""

import openpyxl

from bruv_app.export import export_project
from bruv_app.metrics import Observation, double_count_checks

from .conftest import label_like_perfect_human


def _sheet(wb, name):
    return list(wb[name].iter_rows(values_only=True))


def _status(wb):
    return dict(_sheet(wb, "review_status")[1:])


def test_answers_record_their_scope(project):
    card = project.all_fish()[0]
    project.label_fish(card["id"], "species", species="Lutjanus campechanus")
    dets = {d["id"] for d in project.fish_detections(card["id"])}
    assert project.label_scopes() == {card["id"]: dets}
    project.label_fish(card["id"], "pending")
    assert project.label_scopes() == {card["id"]: set()}


def test_a_label_never_spreads_to_detections_added_after_it(project, tmp_path):
    label_like_perfect_human(project)
    before = export_project(project, tmp_path / "before.xlsx")
    maxn_before = {r[0]: r[1] for r in _sheet(openpyxl.load_workbook(before), "MaxN")[1:]}
    # Some other card's tracklet lands on an answered card without anyone reviewing it again
    # (no review action does this today; the safeguard keeps it that way).
    a, b = [f for f in project.all_fish() if f["status"] == "species"][:2]
    project.db.execute("UPDATE tracklets SET fish_id=? WHERE fish_id=?", (a["id"], b["id"]))
    project.db.execute("DELETE FROM fish WHERE id=?", (b["id"],))
    project.commit()
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "after.xlsx"))
    status = _status(wb)
    moved = len(project.fish_detections(a["id"])) - len(project.label_scopes()[a["id"]])
    assert moved > 0
    assert status["detections on answered cards outside the answer's scope (not counted)"] == moved
    assert status["review"].startswith("provisional")
    maxn_after = {r[0]: r[1] for r in _sheet(wb, "MaxN")[1:]}
    assert all(maxn_after.get(s, 0) <= n for s, n in maxn_before.items())  # nothing counted twice or relabelled
    assert all(r[7] == "provisional" for r in _sheet(wb, "MaxN")[1:])


def test_unanswered_cards_keep_the_counts_provisional(project, tmp_path):
    n = len(project.all_fish())
    status = _status(openpyxl.load_workbook(export_project(project, tmp_path / "a.xlsx")))
    assert status["review"] == f"provisional: {n} cards not answered"


def test_tiny_fish_are_out_of_scope_and_reported(project, tmp_path):
    """Tiny fish aren't part of the review (docs/SPEC.md §15): unanswered, they don't keep
    the counts provisional, but the review status says how many there are."""
    settings = project.load_settings()
    settings.tiny_fish_px = 1e9  # every card is tiny
    project.save_settings(settings)
    status = _status(openpyxl.load_workbook(export_project(project, tmp_path / "a.xlsx")))
    n = len(project.all_fish())
    assert status["review"] == "complete"
    assert status["tiny fish (out of the review's scope, not counted unless answered)"] == f"{n} cards not answered"
    assert status[f"  tiny (< {1e9:g} px, hidden until asked for)"].startswith(f"{n} / 0 / 0 / {n} /")

    label_like_perfect_human(project)
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "b.xlsx"))
    assert _status(wb)["review"] == "complete"
    assert _status(wb)["tiny fish (out of the review's scope, not counted unless answered)"] == "0 cards not answered"
    assert all(r[7] == "complete" for r in _sheet(wb, "MaxN")[1:])
    minutes = _sheet(wb, "review_by_minute")
    assert minutes[0] == ("minute", "clip", "from", "detections", "counted", "not_fish", "unreviewed", "unreviewed_tiny")
    total = project.db.execute("SELECT COUNT(*) FROM detections WHERE tracklet_id IS NOT NULL").fetchone()[0]
    assert sum(r[3] for r in minutes[1:]) == total and sum(r[6] for r in minutes[1:]) == 0


def test_record_rows_follow_their_own_gap_not_the_merge_window(project, tmp_path):
    label_like_perfect_human(project)
    # Pieces of one fish a few seconds apart, each on its own card, all answered alike.
    card = next(f for f in project.all_fish() if len(project.fish_tracklets(f["id"])) > 1)
    for piece in [card["id"], *project.unmerge_fish(card["id"])]:
        project.label_fish(piece, "species", species=card["species"])

    def record_rows(**changes):
        settings = project.load_settings()
        for key, value in changes.items():
            setattr(settings, key, value)
        project.save_settings(settings)
        wb = openpyxl.load_workbook(export_project(project, tmp_path / f"{len(changes)}{sorted(changes)}.xlsx"))
        ws = wb[[n for n in wb.sheetnames if n.startswith("record")][0]]
        return [row for row in ws.iter_rows(min_row=3, values_only=True) if any(row)]

    default = record_rows()
    assert record_rows(gap_seconds=0.0) == default  # tuning the merge leaves the rows alone
    assert len(record_rows(gap_seconds=10.0, sighting_gap_s=0.0)) > len(default)


def test_double_count_checks():
    box = (0.1, 0.1, 0.3, 0.3)
    obs = [
        Observation(1.0, 1, 10, "snapper", box),
        Observation(1.0, 2, 11, "snapper", (0.11, 0.1, 0.31, 0.3)),  # the same fish boxed twice
        Observation(1.0, 3, 12, "jack", (0.12, 0.1, 0.32, 0.3)),  # overlapping, but another label
        Observation(2.0, 4, 10, "snapper", box),
        Observation(2.0, 5, 10, "snapper", (0.6, 0.6, 0.8, 0.8)),  # one card, two boxes at once
        Observation(3.0, 6, 13, "snapper", (0.1, 0.1, 0.5, 0.5)),
        Observation(3.0, 7, 14, "snapper", (0.15, 0.15, 0.25, 0.25)),  # a box inside another
    ]
    checks = [(c.t, c.kind, c.det_ids) for c in double_count_checks(obs)]
    assert checks == [
        (1.0, "overlapping boxes", (1, 2)),
        (2.0, "one card, two boxes", (4, 5)),
        (3.0, "overlapping boxes", (6, 7)),
    ]


def test_export_lists_possible_double_counts(project, tmp_path):
    label_like_perfect_human(project)
    card = next(f for f in project.all_fish() if f["status"] == "species")
    d = project.fish_detections(card["id"])[0]
    project.db.execute(
        "INSERT INTO detections(frame_id, x1, y1, x2, y2, score, tracklet_id) VALUES(?,?,?,?,?,?,?)",
        (d["frame_id"], d["x1"] + 0.003, d["y1"], d["x2"] + 0.003, d["y2"], 0.5, d["tracklet_id"]),
    )
    project.commit()
    project.label_fish(card["id"], "species", species=card["species"])  # reviewed with both boxes
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "out.xlsx"))
    checks = _sheet(wb, "checks")
    assert checks[0] == ("check", "label", "clip", "time", "detections", "cards", "on_maxn_frame")
    assert [r[0] for r in checks[1:]] == ["overlapping boxes"] and checks[1][1] == card["species"]
    assert _status(wb)["double-count checks"].startswith("1 ")
