from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
import pytest

from bruv_app import demo
from bruv_app.config import Settings
from bruv_app.pipeline import Pipeline, create_project
from bruv_app.track import iou
from bruv_app.detect import Detection


@pytest.fixture(scope="session")
def synthetic_videos(tmp_path_factory) -> list[Path]:
    return demo.make_synthetic_videos(tmp_path_factory.mktemp("videos"))


@pytest.fixture
def mock_template(tmp_path) -> Path:
    return demo.make_mock_template(tmp_path / "template" / "viewing template v.2024.xlsx")


# Invented groups and names (not the owner's species list).
GROUPS = {"mullet": ["mullet family", "bluespot mullet"], "goatfish": ["goatfish family", "yellowfin goatfish"]}


@pytest.fixture
def dependent_template(tmp_path) -> Path:
    """A dependent species list, species =INDIRECT(type cell), where each type names a range
    of species, and clip_in/clip_out picked from clip_1… headers on the info sheet through
    one x14 dropdown that covers both columns."""
    import io

    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    dd = wb.active
    dd.title = demo.SHEET_DROPDOWN
    for col, (group, species) in enumerate(GROUPS.items(), start=1):
        letter = get_column_letter(col)
        dd.cell(1, col, group + (" " if group == "goatfish" else ""))  # a stray trailing space
        for row, name in enumerate(species, start=2):
            dd.cell(row, col, name)
        ref = f"'{demo.SHEET_DROPDOWN}'!${letter}$2:${letter}${len(species) + 1}"
        wb.defined_names[group] = DefinedName(group, attr_text=ref)
    wb.defined_names["type"] = DefinedName("type", attr_text=f"'{demo.SHEET_DROPDOWN}'!$A$1:$B$1")
    info = wb.create_sheet(demo.SHEET_INFO)
    info.append(["date", "clip_1", "clip_2", "clip_3"])
    rec = wb.create_sheet(demo.SHEET_RECORD)
    rec.append(["clip_in", "time_in", "clip_out", "time_out", "type", "species", "count"])
    for formula, cells in (("type", "E2:E1000"), ("INDIRECT(E2)", "F2:F1000")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(cells)
        rec.add_data_validation(dv)
    buf = io.BytesIO()
    wb.save(buf)
    clips = [(f"'{demo.SHEET_INFO}'!$B$1:$D$1", "C2:C1000 A2:A1000")]
    return demo.inject_x14_validations(buf.getvalue(), tmp_path / "dependent.xlsx", {demo.SHEET_RECORD: clips})


def demo_settings(**overrides) -> Settings:
    settings = Settings(detector="color", classifier="histogram")
    for key, value in overrides.items():
        setattr(settings, key, value)
    return settings


@pytest.fixture(scope="session")
def processed_project(tmp_path_factory, synthetic_videos):
    """A fully processed synthetic project; tests copy it before changing anything."""
    root = tmp_path_factory.mktemp("base")
    template = demo.make_mock_template(root / "template.xlsx")
    project = create_project(root / "project", synthetic_videos, template, demo_settings())
    Pipeline(project, classifier=demo.DemoClassifier()).run()
    project.db.close()
    return root


@pytest.fixture
def project(processed_project, tmp_path):
    from bruv_app.store import Project

    dst = tmp_path / "copy"
    shutil.copytree(processed_project, dst)
    project = Project(dst / "project")
    # The template path stored in the project points at the session copy; keep it there.
    return project


def ground_truth_species(det) -> str | None:
    """Species of the scene fish best matching a stored detection (the 'perfect human')."""
    box = Detection(det["x1"], det["y1"], det["x2"], det["y2"], 1.0)
    best, best_iou = None, 0.0
    for species, (x1, y1, x2, y2) in demo.scene_at(det["t"]):
        o = iou(box, Detection(x1, y1, x2, y2, 1.0))
        if o > best_iou:
            best, best_iou = species, o
    return best if best_iou > 0.3 else None


def label_like_perfect_human(project) -> None:
    for f in project.all_fish():
        votes = [ground_truth_species(d) for d in project.fish_detections(f["id"])]
        votes = [v for v in votes if v]
        if not votes:
            project.label_fish(f["id"], "not_fish")
            continue
        species = max(set(votes), key=votes.count)
        project.label_fish(f["id"], "species", species=species)


def expected_maxn(fps: float = 1.0) -> dict[str, int]:
    total = demo.CHAPTER_S * 2
    best: dict[str, int] = {}
    for k in range(int(total * fps)):
        visible = [s for s, _ in demo.scene_at(k / fps)]
        for species in set(visible):
            best[species] = max(best.get(species, 0), visible.count(species))
    return best


@pytest.fixture
def rng():
    return np.random.default_rng(0)
