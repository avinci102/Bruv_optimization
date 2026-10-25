"""Families: the template's groups, scores added up per family, and MaxN per family."""

import numpy as np
import openpyxl
import pytest
from fastapi.testclient import TestClient

from bruv_app import demo
from bruv_app.export import export_project
from bruv_app.families import extra_lines, family_of, rank_families
from bruv_app.pipeline import Pipeline, create_project
from bruv_app.server import create_app

from .conftest import demo_settings, ground_truth_species


def test_extra_list_lines_can_name_a_family(tmp_path):
    (tmp_path / "extra_species.txt").write_text(
        "# my reef\nunicornfish > surgeonfish\nsurgeonfish\n\nangelfish\nangelfish\n", encoding="utf-8"
    )
    assert extra_lines(tmp_path / "template.xlsx") == [
        ("unicornfish", "surgeonfish"), ("surgeonfish", None), ("angelfish", None),
    ]


@pytest.fixture
def grouped(tmp_path, synthetic_videos, dependent_template):
    project = create_project(tmp_path / "p", synthetic_videos, dependent_template, demo_settings(), ["batfish"])
    Pipeline(project, classifier=demo.DemoClassifier()).run()
    return project


def test_families_come_from_the_templates_groups(grouped):
    names = ["mullet family", "bluespot mullet", "goatfish family", "yellowfin goatfish",
             "other (indicate in comment)", "batfish", "nemo"]
    assert family_of(grouped, names) == {
        "mullet family": "mullet", "bluespot mullet": "mullet",
        "goatfish family": "goatfish", "yellowfin goatfish": "goatfish",  # the template's stray space is trimmed
        "other (indicate in comment)": "other (indicate in comment)",  # a mixed list ("interest") isn't a family
        "batfish": "batfish", "nemo": "nemo",  # an extra name, or any other name, is its own family
    }
    assert grouped.get_meta("species_groups")["bluespot mullet"] == "mullet"  # read once, then kept


def test_mixed_lists_are_not_families():
    from bruv_app.families import is_family

    assert is_family("grouper", ["grouper family", "blue-and-yellow grouper", "giant grouper"])
    assert is_family("shark", ["blacktip reef shark", "great hammerhead", "whale shark"])  # most carry the name
    assert is_family("turtle", ["green sea turtle", "unknown sea turtle"])
    assert is_family("snapper", ["snapper family", "chinamanfish", "green jobfish"])  # has a family entry
    assert not is_family("spec_of_interest", ["bigfin reef squid", "black marlin", "napoleon wrasse"])


def test_families_rank_by_their_best_name():
    """Adding up would put mullet first (0.5 against 0.4); ranked by the best name, the
    goatfish stays first, as it does by species, so a long list can't win by its length."""
    labels = ["mullet family", "bluespot mullet", "goatfish family", "batfish"]
    family = {"mullet family": "mullet", "bluespot mullet": "mullet", "goatfish family": "goatfish", "batfish": "batfish"}
    ranked = rank_families(labels, np.array([0.2, 0.3, 0.4, 0.1]), family)
    assert [(f.key, round(f.best, 2), round(f.p, 2)) for f in ranked] == [
        ("goatfish", 0.4, 0.4), ("mullet", 0.3, 0.5), ("batfish", 0.1, 0.1)]
    mullet = ranked[1]
    assert mullet.name == "mullet" and mullet.entry == "mullet family"
    assert [(n, round(p, 2)) for n, p in mullet.members] == [("bluespot mullet", 0.3), ("mullet family", 0.2)]
    assert ranked[2].entry == "batfish"  # a one-name family is recorded by its name


def test_card_detail_lists_the_top_families(grouped):
    client = TestClient(create_app(grouped.root))
    card = grouped.all_fish()[0]["id"]
    families = client.get(f"/api/fish/{card}").json()["families"]
    assert 1 <= len(families) <= 3
    assert [f["best"] for f in families] == sorted((f["best"] for f in families), reverse=True)
    species = client.get(f"/api/fish/{card}").json()["guesses"]
    assert families[0]["members"][0][0] == species[0][0]  # the top family holds the top species
    for f in families:
        assert f["best"] == f["members"][0][1] and f["p"] >= f["best"]
    by_key = {f["key"]: f for f in families}
    if "mullet" in by_key:
        assert by_key["mullet"]["entry"] == "mullet family" and by_key["mullet"]["name"] == "mullet"


def test_maxn_by_family_counts_each_frame_once(grouped, tmp_path):
    """Two of the yellow school answered as different names of one family count together,
    frame by frame; the species' own MaxN rows stay as they are."""
    at_17 = [d for f in grouped.all_fish() for d in grouped.fish_detections(f["id"]) if d["t"] == 17.0]
    school = sorted({grouped.db.execute("SELECT fish_id FROM tracklets WHERE id=?", (d["tracklet_id"],)).fetchone()[0]
                     for d in at_17 if ground_truth_species(d) == demo.YELLOW})
    assert len(school) == 3  # fish on screen together are always on separate cards
    for card, name in zip(school, ["mullet family", "bluespot mullet", "goatfish family"]):
        grouped.label_fish(card, "species", species=name)
    wb = openpyxl.load_workbook(export_project(grouped, tmp_path / "out.xlsx"))
    families = {r[0]: r for r in wb["MaxN_by_family"].iter_rows(min_row=2, values_only=True)}
    assert families["mullet"][1] == 2 and families["mullet"][6] == "bluespot mullet; mullet family"
    assert families["goatfish"][1] == 1
    species = {r[0]: r[1] for r in wb["MaxN"].iter_rows(min_row=2, values_only=True)}
    assert species["mullet family"] == 1 and species["bluespot mullet"] == 1
