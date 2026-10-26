"""The pilot's review workflow (docs/SPEC.md §15): best views, two-level answers, the
look-alike list's log, the review window, and the setup check."""

import openpyxl
import pytest
from fastapi.testclient import TestClient

from bruv_app import demo
from bruv_app.cli import main, parse_clock
from bruv_app.export import export_project
from bruv_app.families import resolve_answer
from bruv_app.server import create_app
from bruv_app.store import Project
from bruv_app.views import best_view, view_score

from .conftest import label_like_perfect_human


def _det(x1, y1, x2, y2, quality=1.0):
    return {"id": (x1, y1), "x1": x1, "y1": y1, "x2": x2, "y2": y2, "quality": quality, "width": 960, "height": 540}


def test_the_best_view_is_side_on_and_whole():
    side = _det(0.30, 0.40, 0.50, 0.46)  # 192 x 32 px: a fish lying across
    head = _det(0.30, 0.40, 0.34, 0.48)  # 38 x 43 px: about as tall as long, head-on
    cut = _det(0.00, 0.40, 0.20, 0.46)  # side-on, but touching the frame's edge
    assert view_score(side) > view_score(head) and view_score(side) > view_score(cut)
    assert best_view([head, cut, side]) is side and best_view([]) is None
    # Quality (sharpness, size, confidence) still counts: a much sharper head-on view wins.
    assert best_view([side, _det(0.30, 0.40, 0.34, 0.48, quality=5.0)])["quality"] == 5.0


def _grouped(project):
    project.set_meta("species_groups", {
        "grouper family": "grouper", "giant grouper": "grouper",
        "giant moray eel": "moray_eel", "snowflake moray eel": "moray_eel",  # no family entry
        "unicornfish": "unicornfish",
    })
    return ["grouper family", "giant grouper", "giant moray eel", "snowflake moray eel", "unicornfish"]


def test_the_family_is_required_and_the_species_optional(tmp_path):
    project = Project(tmp_path / "p")
    names = _grouped(project)
    assert resolve_answer(project, names, "giant grouper", None) == ("giant grouper", "grouper")
    assert resolve_answer(project, names, "giant grouper", "grouper") == ("giant grouper", "grouper")
    assert resolve_answer(project, names, None, "grouper") == ("grouper family", "grouper")  # species unknown
    assert resolve_answer(project, names, None, "moray_eel") == ("moray_eel", "moray_eel")  # no family entry
    assert resolve_answer(project, names, "moray_eel", None) == ("moray_eel", "moray_eel")  # as repeated by R
    assert resolve_answer(project, names, "unicornfish", None) == ("unicornfish", "unicornfish")
    for species, family in (("giant grouper", "moray_eel"), ("nemo", None), (None, None), (None, "shark")):
        with pytest.raises(ValueError):
            resolve_answer(project, names, species, family)


def test_a_family_without_an_entry_is_recorded_and_exported(project, tmp_path):
    dropdowns = project.get_meta("dropdowns")
    dropdowns["species"] = dropdowns["species"] + ["giant moray eel", "snowflake moray eel"]
    project.set_meta("dropdowns", dropdowns)
    project.set_meta("species_groups", {"giant moray eel": "moray_eel", "snowflake moray eel": "moray_eel"})
    client = TestClient(create_app(project.root))
    card = project.all_fish()[0]["id"]
    answered = client.post(f"/api/fish/{card}/label", json={"status": "species", "family": "moray_eel", "confirm": True})
    assert answered.status_code == 200, answered.text
    assert project.fish(card)["species"] == "moray_eel" and project.answer_family(card) == "moray_eel"
    assert client.get(f"/api/fish/{card}").json()["family"] == "moray_eel"
    bad = client.post(f"/api/fish/{card}/label", json={"status": "species", "family": "shark"})
    assert bad.status_code == 400
    wb = openpyxl.load_workbook(export_project(project, tmp_path / "out.xlsx"))
    cells = [v for row in wb[demo.SHEET_RECORD].iter_rows(values_only=True) for v in row if isinstance(v, str)]
    assert any(v.startswith("moray eel") for v in cells)  # the "other" entry's comment, with spaces


def test_the_look_alike_list_logs_what_was_left_unticked(project):
    client = TestClient(create_app(project.root))
    source = next(f["id"] for f in project.all_fish()
                  if len(client.get(f"/api/fish/{f['id']}/lookalikes").json()["cards"]) >= 2)
    cards = client.get(f"/api/fish/{source}/lookalikes").json()["cards"]
    client.post(f"/api/fish/{source}/label", json={"status": "species", "species": demo.RED, "confirm": True})
    first, second = cards[0]["id"], cards[1]["id"]
    res = client.post("/api/lookalikes/answer", json={"source": source, "ids": [first], "status": "species",
                                                      "species": demo.RED, "shown": [first, second]}).json()
    log = {r["fish_id"]: dict(r) for r in project.db.execute("SELECT * FROM suggestion_log")}
    assert log[second]["decision"] == "unticked" and log[second]["source"] == source
    assert log[first]["decision"] == ("answered" if first in res["answered"] else "left")
    assert log[first]["similarity"] == pytest.approx(cards[0]["similarity"], abs=1e-3)


def test_the_review_window_hides_and_uncounts_setup_time(project, tmp_path):
    label_like_perfect_human(project)
    status = lambda wb: {r[0]: r[1] for r in wb["review_status"].iter_rows(values_only=True)}  # noqa: E731
    before = status(openpyxl.load_workbook(export_project(project, tmp_path / "before.xlsx")))
    first = min(project.all_fish(), key=lambda f: f["start_t"])
    project.set_review_window(first["end_t"] + 0.5, None)
    outside = project.outside_window()
    assert first["id"] in outside
    client = TestClient(create_app(project.root))
    listed = {c["id"]: c for c in client.get("/api/fish").json()}
    assert listed[first["id"]]["outside"] is True and not listed[max(listed)]["outside"]
    assert client.get("/api/summary").json()["outside_window"] == len(outside)
    after = status(openpyxl.load_workbook(export_project(project, tmp_path / "after.xlsx")))
    assert after["detections counted"] < before["detections counted"]
    assert any(k.startswith("review window") for k in after)
    with pytest.raises(ValueError):
        project.set_review_window(10.0, 5.0)


def test_the_review_window_command(project, capsys):
    main(["review-window", "--project", str(project.root), "--start", "0:05", "--end", "0:20"])
    assert project.review_window() == (5.0, 20.0)
    assert "Review window: 00:00:05 to 00:00:20" in capsys.readouterr().out
    main(["review-window", "--project", str(project.root), "--clear"])
    assert project.review_window() == (None, None)
    assert parse_clock("49") == 49 and parse_clock("1:02:05") == 3725 and parse_clock(None) is None


def test_the_setup_check_says_where_processing_runs(capsys):
    main(["check-setup"])
    out = capsys.readouterr().out
    assert "Processing will run on:" in out and "NVIDIA graphics card:" in out
