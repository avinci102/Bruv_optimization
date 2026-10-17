"""Read dropdown lists from the owner's Excel template and write exports into a copy of it.

Rules: write by header name (never by column letter), never touch the original file, and
take dropdown lists from the sheet's data-validation rules.

Excel stores dropdowns that point at another sheet in an extension block (x14) that
openpyxl drops on load (it only warns). So validations are read from the raw sheet XML,
and any that openpyxl dropped are re-attached, with their settings, to the exported copy.
"""

from __future__ import annotations

import io
import logging
import posixpath
import re
import warnings
import zipfile
from copy import copy
from dataclasses import dataclass, field
from datetime import time as dtime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

log = logging.getLogger(__name__)

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
    "xm": "http://schemas.microsoft.com/office/excel/2006/main",
}

RECORD_PREFIX = "record"
DROPDOWN_PREFIX = "dropdown"
INFO_PREFIX = "info"
TIME_HEADERS = {"time_in", "time_out"}


@dataclass
class RawValidation:
    type: str
    formula1: str
    sqref: str
    from_extension: bool
    allow_blank: bool = True
    # Error/prompt settings (showErrorMessage, errorStyle, prompt, ...) carried to the copy.
    options: dict = field(default_factory=dict)


@dataclass
class TemplateInfo:
    path: Path
    record_sheet: str
    dropdown_sheet: str | None
    header_row: int
    headers: dict[str, int]  # normalised header -> 1-based column index
    dropdowns: dict[str, list[str]] = field(default_factory=dict)
    validations: list[RawValidation] = field(default_factory=list)
    next_row: int = 0  # first record row after any existing data (example row included)
    # Dependent dropdowns, e.g. species =INDIRECT(E2) where column E is a species group:
    # child header -> (parent header, {child option: the parent option it is listed under}).
    dependent: dict[str, tuple[str, dict[str, str]]] = field(default_factory=dict)

    def options(self, header: str) -> list[str]:
        return self.dropdowns.get(norm(header), [])


def norm(header) -> str:
    return str(header).strip().lower() if header is not None else ""


def _find_sheet(names: list[str], prefix: str) -> str | None:
    for name in names:
        if name.strip().lower().startswith(prefix):
            return name
    return None


# --- raw XML access -------------------------------------------------------------------


def sheet_xml_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    """Sheet name -> path of its XML part inside the xlsx zip."""
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels.findall("pr:Relationship", NS)}
    out = {}
    for sheet in wb.findall("m:sheets/m:sheet", NS):
        target = targets[sheet.get(f"{{{NS['r']}}}id")]
        path = target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
        out[sheet.get("name")] = posixpath.normpath(path)
    return out


def _external_link_rels(path: Path) -> list[dict]:
    """For each external link, in workbook order: its relationships by id."""
    with zipfile.ZipFile(path) as zf:
        wb = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        targets = {r.get("Id"): r.get("Target") for r in rels.findall("pr:Relationship", NS)}
        out = []
        for ref in wb.findall("m:externalReferences/m:externalReference", NS):
            target = targets[ref.get(f"{{{NS['r']}}}id")]
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            link_rels = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
            found = ET.fromstring(zf.read(link_rels)) if link_rels in zf.namelist() else None
            out.append({} if found is None else {r.get("Id"): r for r in found.findall("pr:Relationship", NS)})
    return out


def _defined_names(zf: zipfile.ZipFile) -> dict[str, str]:
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    return {
        d.get("name").lower(): (d.text or "").strip()
        for d in wb.findall("m:definedNames/m:definedName", NS)
    }


_BOOL_OPTIONS = ("showErrorMessage", "showInputMessage", "showDropDown")
_TEXT_OPTIONS = ("errorStyle", "promptTitle", "prompt", "errorTitle", "error", "operator")


def _options(el) -> dict:
    opts = {k: el.get(k) in ("1", "true") for k in _BOOL_OPTIONS if el.get(k) is not None}
    opts.update({k: el.get(k) for k in _TEXT_OPTIONS if el.get(k) is not None})
    return opts


def read_raw_validations(path: Path, sheet_name: str) -> list[RawValidation]:
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read(sheet_xml_paths(zf)[sheet_name]))
    out = []
    for dv in root.findall("m:dataValidations/m:dataValidation", NS):
        f1 = dv.find("m:formula1", NS)
        out.append(
            RawValidation(
                dv.get("type", ""),
                (f1.text or "").strip() if f1 is not None else "",
                dv.get("sqref", ""),
                False,
                dv.get("allowBlank", "0") in ("1", "true"),
                _options(dv),
            )
        )
    for dv in root.iter(f"{{{NS['x14']}}}dataValidation"):
        f = dv.find("x14:formula1/xm:f", NS)
        sq = dv.find("xm:sqref", NS)
        out.append(
            RawValidation(
                dv.get("type", ""),
                (f.text or "").strip() if f is not None else "",
                (sq.text or "").strip() if sq is not None else "",
                True,
                dv.get("allowBlank", "0") in ("1", "true"),
                _options(dv),
            )
        )
    return out


# --- reading ----------------------------------------------------------------------------

_REF = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^'!]+)!)?(?P<range>\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)$")
# A dependent list: =INDIRECT(E2) shows the range named by whatever was picked in E2.
_INDIRECT = re.compile(r"^INDIRECT\(\s*\$?(?P<col>[A-Z]{1,3})\$?\d+\s*\)$", re.IGNORECASE)


def _unique(values) -> list[str]:
    seen, out = set(), []
    for v in values:
        if v is None:
            continue
        s = str(v).strip()
        if s and s not in seen:
            seen.add(s)
            out.append(s)
    return out


def _resolve_formula(formula: str, wb, default_sheet: str, names: dict[str, str]) -> list[str] | None:
    f = formula.strip().lstrip("=")
    if not f:
        return None
    if f.startswith('"') and f.endswith('"'):
        return _unique(f[1:-1].split(","))
    if f.lower() in names:
        return _resolve_formula(names[f.lower()], wb, default_sheet, names)
    m = _REF.match(f)
    if not m:
        return None  # INDIRECT/OFFSET and other dynamic lists aren't supported
    sheet = m.group("sheet") or default_sheet
    if sheet.startswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    if sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    min_col, min_row, max_col, max_row = range_boundaries(m.group("range").replace("$", ""))
    values = []
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        values.extend(c.value for c in row)
    return _unique(values)


def _sqref_columns(sqref: str) -> set[int]:
    cols = set()
    for part in sqref.split():
        min_col, _, max_col, _ = range_boundaries(part.replace("$", ""))
        cols.update(range(min_col, (max_col or min_col) + 1))
    return cols


def _load(path: Path, data_only: bool = False, report: bool = False, rich_text: bool = False):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(
            path, data_only=data_only, keep_vba=path.suffix.lower() == ".xlsm", rich_text=rich_text
        )
    if report:
        for w in caught:
            message = str(w.message)
            if "Data Validation extension" in message:
                continue  # re-attached by export_workbook
            log.warning("Template feature openpyxl can't keep in the export: %s", message)
    return wb


def read_template(path: Path) -> TemplateInfo:
    path = Path(path)
    wb = _load(path, data_only=True)
    record = _find_sheet(wb.sheetnames, RECORD_PREFIX)
    if record is None:
        raise ValueError(f"No sheet starting with '{RECORD_PREFIX}' in {path.name}")
    dropdown = _find_sheet(wb.sheetnames, DROPDOWN_PREFIX)
    ws = wb[record]

    header_row, headers = 1, {}
    for r in range(1, min(ws.max_row, 10) + 1):
        row = {norm(c.value): c.column for c in ws[r] if c.value is not None}
        if "species" in row and ("time_in" in row or "clip_in" in row):
            header_row, headers = r, row
            break
    if not headers:
        raise ValueError(f"Couldn't find the header row (species, time_in, ...) in '{record}'")

    with zipfile.ZipFile(path) as zf:
        names = _defined_names(zf)
    validations = read_raw_validations(path, record)
    col_to_header = {c: h for h, c in headers.items()}
    dropdowns: dict[str, list[str]] = {}
    dependent_cols: list[tuple[set[int], int]] = []  # (child columns, parent column)
    for v in validations:
        if v.type != "list":
            continue
        if m := _INDIRECT.match(v.formula1.strip().lstrip("=")):
            parent_col = column_index_from_string(m.group("col").upper())
            dependent_cols.append((_sqref_columns(v.sqref), parent_col))
            continue
        values = _resolve_formula(v.formula1, wb, record, names)
        if not values:
            continue
        for col in _sqref_columns(v.sqref):
            if col in col_to_header and col_to_header[col] not in dropdowns:
                dropdowns[col_to_header[col]] = values

    # Fallback: a column on the dropdown sheet whose header matches.
    if dropdown:
        dws = wb[dropdown]
        for col in dws.iter_cols(min_row=1, max_row=dws.max_row):
            if not col or col[0].value is None:
                continue
            h = norm(col[0].value)
            if h in headers and h not in dropdowns:
                values = _unique(c.value for c in col[1:])
                if values:
                    dropdowns[h] = values

    # Dependent lists: each option of the parent's list (e.g. a species group) names a range
    # of child options (species). The child list is all of them, and each child option
    # remembers its parent, so the export can fill the parent column to match.
    dependent: dict[str, tuple[str, dict[str, str]]] = {}
    for child_cols, parent_col in dependent_cols:
        parent = col_to_header.get(parent_col)
        groups: dict[str, str] = {}
        for option in dropdowns.get(parent, []):
            for value in _resolve_formula(option, wb, record, names) or []:
                groups.setdefault(value, option)
        for col in child_cols:
            child = col_to_header.get(col)
            if groups and child and child not in dropdowns:
                dropdowns[child] = list(groups)
                dependent[child] = (parent, groups)

    next_row = _first_empty_row(ws, header_row, list(headers.values()))
    return TemplateInfo(
        path, record, dropdown, header_row, headers, dropdowns, validations, next_row, dependent
    )


# --- writing ----------------------------------------------------------------------------


def seconds_to_time(seconds: float) -> dtime:
    s = int(round(max(0.0, seconds)))
    return dtime((s // 3600) % 24, (s % 3600) // 60, s % 60)


def _first_empty_row(ws, header_row: int, columns: list[int]) -> int:
    last = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in columns):
            last = r
    return last + 1


def _write_table(wb, title: str, header: list[str], rows: list[list]) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(str(h)) + 4))
    ws.freeze_panes = "A2"


_SHAPE = re.compile(rb"<(?:\w+:)?sp[\s>]")


def _warn_about_shapes(path: Path) -> None:
    """openpyxl drops shapes and text boxes without a warning, so look for them ourselves."""
    with zipfile.ZipFile(path) as zf:
        parts = [n for n in zf.namelist() if n.startswith("xl/drawings/") and n.endswith(".xml")]
        with_shapes = [n for n in parts if _SHAPE.search(zf.read(n))]
    if with_shapes:
        log.warning(
            "The template has shapes or text boxes (%d drawing part(s)); they can't be kept "
            "in the export, so check the copy against the template.",
            len(with_shapes),
        )


def _with_parents(info: TemplateInfo, row: dict) -> dict:
    """Set the parent of each dependent dropdown from its child (type from species), so the
    child's value is in the list Excel shows for that row."""
    out = dict(row)
    keys = {norm(k): k for k in row}
    for child, (parent, groups) in info.dependent.items():
        value = row.get(keys.get(child, child))
        if value in groups:
            out[keys.get(parent, parent)] = groups[value]
    return out


def export_workbook(
    info: TemplateInfo,
    out_path: Path,
    record_rows: list[dict],
    maxn_header: list[str],
    maxn_rows: list[list],
    details_header: list[str],
    details_rows: list[list],
) -> dict:
    """Write record rows (by header name) plus MaxN and app_details sheets into a copy.

    Returns {"first_row": int, "rows": int, "skipped_headers": [...]}.
    """
    out_path = Path(out_path)
    if out_path.resolve() == Path(info.path).resolve():
        raise ValueError("Refusing to overwrite the original template; choose another file.")
    # rich_text keeps bold/italic words inside a cell (e.g. on a readme sheet).
    wb = _load(Path(info.path), report=True, rich_text=True)
    _warn_about_shapes(Path(info.path))
    ws = wb[info.record_sheet]

    # openpyxl keeps only the first relationship of an external link, but the link points
    # at its own (a link saved with alternate paths lists several). Excel refuses the whole
    # file as damaged when that relationship is missing, so put it back.
    for link, rels in zip(wb._external_links, _external_link_rels(Path(info.path))):
        book, kept = link.externalBook, link.file_link
        if book is not None and kept is not None and book.id in rels and kept.Id != book.id:
            kept.Id, kept.Target = book.id, rels[book.id].get("Target")
            kept.TargetMode = rels[book.id].get("TargetMode")

    # Re-attach dropdowns openpyxl dropped (x14 extension, i.e. cross-sheet lists), on
    # every sheet, not only `record` (the info sheet may have them too).
    for sheet in wb.sheetnames:
        if sheet in ("MaxN", "app_details"):
            continue
        try:
            raw = read_raw_validations(Path(info.path), sheet)
        except KeyError:  # a sheet the template doesn't have
            continue
        for v in raw:
            if v.from_extension and v.formula1 and v.sqref:
                # sqref may hold several ranges ("C2:C1000 A2:A1000"); the constructor
                # parses them all, unlike dv.add().
                dv = DataValidation(
                    type=v.type or "list", formula1=v.formula1, allow_blank=v.allow_blank,
                    sqref=v.sqref, **v.options,
                )
                wb[sheet].add_data_validation(dv)

    start = info.next_row
    example_row = info.header_row + 1
    skipped: set[str] = set()
    for i, row in enumerate(record_rows):
        r = start + i
        for header, value in _with_parents(info, row).items():
            col = info.headers.get(norm(header))
            if col is None:
                skipped.add(header)
                continue
            if value is None or value == "":
                continue
            cell = ws.cell(r, col)
            if not cell.has_style:
                # Excel shows an untouched cell with its row's or column's format (e.g. a
                # centred count column); openpyxl gives a new cell none, so copy it.
                row_dim = ws.row_dimensions[r]
                dim = row_dim if row_dim.customFormat else ws.column_dimensions[get_column_letter(col)]
                cell._style = copy(dim._style)
            cell.value = value
            fmt = ws.cell(example_row, col).number_format
            if norm(header) in TIME_HEADERS:
                cell.number_format = fmt if fmt and fmt != "General" else "hh:mm:ss"

    _write_table(wb, "MaxN", maxn_header, maxn_rows)
    _write_table(wb, "app_details", details_header, details_rows)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    _save(wb, out_path)
    return {"first_row": start, "rows": len(record_rows), "skipped_headers": sorted(skipped)}


_BLANK_TEXT = re.compile(rb"<t>(\s+)</t>")


def _save(wb, out_path: Path) -> None:
    """Save, then mark whitespace-only text runs xml:space="preserve". openpyxl leaves it
    off (e.g. a rich-text run holding one space), Excel reads the run as empty, and then
    refuses to open the whole file as damaged."""
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/") or item.filename == "xl/sharedStrings.xml":
                data = _BLANK_TEXT.sub(rb'<t xml:space="preserve">\1</t>', data)
            zout.writestr(item, data)

